from copy import deepcopy
import hashlib
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt


template_path = Path(r"C:\Users\SHELL\Downloads\entregable de evidencias.docx")
captures_path = Path(r"C:\Users\SHELL\Downloads\capturas.docx")
output_path = Path(r"C:\Users\SHELL\Downloads\MatriculaBackend\evidencias\Entregable_Evidencias_EduAndes.docx")
output_path.parent.mkdir(parents=True, exist_ok=True)

expected_template_sha256 = "220bb188f8203a7cfe4ffd90f019f6e091c50a11cdc58f591e99ddbbe113f6d2"
actual_template_sha256 = hashlib.sha256(template_path.read_bytes()).hexdigest()
if actual_template_sha256 != expected_template_sha256:
    raise RuntimeError("La portada de referencia cambió desde la revisión visual.")

document = Document(template_path)

for paragraph in document.paragraphs:
    if paragraph.text.startswith("Laboratorio: Registro transaccional"):
        paragraph.runs[0].text = "Evidencias del caso EduAndes Matrícula Académica"
        for run in paragraph.runs[1:]:
            run.text = ""
    elif paragraph.text == "Lima, Agosto 2026":
        paragraph.runs[0].text = "Lima, septiembre de 2026"
        for run in paragraph.runs[1:]:
            run.text = ""

# Retain the cover and its student-name table; remove everything after its date.
body = document._element.body
date_paragraph = next(p for p in document.paragraphs if p.text == "Lima, septiembre de 2026")
after_date = False
for child in list(body):
    if child is date_paragraph._p:
        after_date = True
        continue
    if after_date and child.tag != qn("w:sectPr"):
        body.remove(child)

# Start evidence on a new page and match the source screenshot document's image area.
evidence_section = document.add_section(WD_SECTION_START.NEW_PAGE)
evidence_section.page_width = Mm(210)
evidence_section.page_height = Mm(297)
evidence_section.top_margin = Inches(0.5)
evidence_section.bottom_margin = Inches(0.5)
evidence_section.left_margin = Inches(0.5)
evidence_section.right_margin = Inches(0.5)

captures = Document(captures_path)
caption_updates = {
    "RN-03: no mas de una matrícula en estado REGISTRADA en el mismo periodo.":
        "RN-03: una matrícula REGISTRADA por estudiante y periodo",
    "Salud:": "Salud de la API:",
    "BUSQUEDA:": "Búsqueda de cursos:",
    "Matricula:": "Matrícula correcta:",
    "Anulada:": "Anulación de matrícula:",
}
evidence_titles = {
    "RN-01: estudiante inactivo (Carlos)",
    "RN-01: curso de otra carrera (María es de Civil; IS401 es de Sistemas)",
    "RN-02: curso sin vacantes (IS404)",
    "RN-03: una matrícula REGISTRADA por estudiante y periodo",
    "RN-04: más de 20 créditos",
    "Salud de la API:",
    "Búsqueda de cursos:",
    "Matrícula correcta:",
    "Anulación de matrícula:",
}
first_evidence_title = True

for source_paragraph in captures.paragraphs:
    replacement = caption_updates.get(source_paragraph.text)
    if replacement and source_paragraph.runs:
        source_paragraph.runs[0].text = replacement
        for run in source_paragraph.runs[1:]:
            run.text = ""

    if source_paragraph.text in evidence_titles:
        if not first_evidence_title:
            source_paragraph.paragraph_format.page_break_before = True
        first_evidence_title = False
        source_paragraph.paragraph_format.keep_with_next = True
        source_paragraph.paragraph_format.space_after = Pt(8)
        if source_paragraph.runs:
            source_paragraph.runs[0].bold = True
            source_paragraph.runs[0].font.size = Pt(13)

    paragraph_copy = deepcopy(source_paragraph._p)
    for blip in paragraph_copy.xpath(".//a:blip"):
        old_rid = blip.get(qn("r:embed"))
        image_part = captures.part.related_parts[old_rid]
        new_rid = document.part.get_or_add_image(BytesIO(image_part.blob))[0]
        blip.set(qn("r:embed"), new_rid)

    body.insert(len(body) - 1, paragraph_copy)

document.core_properties.title = "Evidencias del caso EduAndes Matrícula Académica"
document.core_properties.subject = "Capturas de pruebas funcionales de MatriculaBackend"
document.save(output_path)
print(output_path)
